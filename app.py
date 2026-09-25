"""
AEON Device Management — Laptops & SIM Card tracker.

Env vars:
  DATABASE_URL   required, e.g. postgresql://root:pwd@host:5432/zeabur
  PORT           optional, default 8080
  ADMIN_PASSWORD gates the /config page (DB connection settings — separate from accounts below)
  SYNC_TOKEN     shared secret for POST /api/device-management/sync (Power Automate)
  BOOTSTRAP_ADMIN_EMAIL     seeds one admin account on first run (default itdept@aeondelightasia.com)
  BOOTSTRAP_ADMIN_PASSWORD  password for that seeded account — required for the seed to happen
  EMBEDDED_IN_IFRAME        set truthy when embedded in OneView's iframe — needed for the
                            session cookie (login, DB config) to survive third-party framing
"""
import json
import os
import logging
from datetime import datetime, timezone

from flask import Flask, render_template, jsonify, request, session, redirect, url_for

import db as dbmod         # engine-aware connection layer (db.py)
import accounts as acctmod  # app-level accounts + business-unit permission control

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger(__name__)

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET", "dev-secret-change-me")

# When embedded inside OneView's iframe (aeon-oneview.zeabur.app), browsers
# treat this app as a third-party site and block the session cookie unless
# SameSite=None + Secure are set. Default OFF so local http://localhost dev
# keeps working. Zeabur prod: set EMBEDDED_IN_IFRAME=1 in Variables.
_iframe_flag = (os.environ.get("EMBEDDED_IN_IFRAME") or "").strip().lower()
if _iframe_flag in ("1", "true", "yes", "on", "y", "t"):
    app.config["SESSION_COOKIE_SAMESITE"] = "None"
    app.config["SESSION_COOKIE_SECURE"]   = True
    log.info("Session cookies: SameSite=None; Secure (iframe-embed mode)")
else:
    log.info(f"Session cookies: default SameSite=Lax (EMBEDDED_IN_IFRAME={_iframe_flag!r})")

ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin")
SYNC_TOKEN     = os.environ.get("SYNC_TOKEN", "")

query = dbmod.query


# ── Site derivation (must mirror device_management.js categoryToSite/simSite,
#    so server-side filtering matches what the site pills already show) ─────

_SITE_ALIASES = {"MC6": "BO"}   # SIM "MC6" == laptop "BO" (same site, different code)


def _laptop_site(category):
    if not category:
        return None
    cat = str(category).strip()
    if cat == "ADG-HK-BO":
        return "BO"
    if cat.startswith("ADG-HK-Site-"):
        return cat[len("ADG-HK-Site-"):]
    if cat.startswith("ADG-HK-"):
        return cat[len("ADG-HK-"):]
    return cat


def _sim_site(contract_site):
    if not contract_site:
        return None
    s = str(contract_site).strip()
    return _SITE_ALIASES.get(s, s)


# ── Device Management (read) ────────────────────────────────────────

def fetch_device_management(allowed_sites=None):
    """allowed_sites=None means unrestricted (admin). Otherwise only rows
    whose derived site is in that set are returned — enforced here, not just
    hidden client-side, so a scoped account can't read other sites via the API."""
    laptops = query("""
        SELECT serial_number, device_name, unique_device, enrollment_date,
               manufacturer, model, category, primary_user_upn,
               primary_user_display_name, device_per_user, remarks, status
        FROM laptops
        ORDER BY primary_user_display_name NULLS LAST, serial_number
    """)
    # `count_day_of_sim_contract_end` is computed live (contract_end_date - today)
    # so the "Days" column reflects current age without needing a re-import.
    sims = query("""
        SELECT sim_no, telcom, contract_site, contract_no, monthly_fee,
               sim_owner, contract_start_date, contract_end_date,
               site_contract,
               CASE
                   WHEN contract_end_date IS NULL THEN NULL
                   ELSE (contract_end_date - CURRENT_DATE)::INTEGER
               END AS count_day_of_sim_contract_end
        FROM sim_cards
        ORDER BY
            CASE WHEN contract_end_date IS NULL THEN 1 ELSE 0 END,
            contract_end_date ASC,
            sim_owner
    """)
    sync = query("SELECT last_synced, laptops_count, sims_count FROM device_mgmt_sync WHERE id=1", fetchone=True)

    def serialise(row):
        out = {}
        for k, v in row.items():
            if hasattr(v, "isoformat"):
                out[k] = v.isoformat()
            elif v is None:
                out[k] = None
            else:
                # Convert Decimal / other non-JSON-native types to string
                try:
                    json.dumps(v)
                    out[k] = v
                except (TypeError, ValueError):
                    out[k] = str(v)
        return out

    laptops = [serialise(r) for r in laptops]
    sims = [serialise(r) for r in sims]

    if allowed_sites is not None:
        laptops = [r for r in laptops if _laptop_site(r["category"]) in allowed_sites]
        sims = [r for r in sims if _sim_site(r["contract_site"]) in allowed_sites]

    return {
        "laptops":  laptops,
        "sims":     sims,
        "summary": {
            "total_laptops":   len(laptops),
            "spared_laptops":  sum(1 for r in laptops if (r["primary_user_display_name"] or "").strip().lower() in ("", "spare", "spared", "-")),
            "total_sims":      len(sims),
            "expiring_30d":    sum(1 for r in sims if r["count_day_of_sim_contract_end"] is not None and r["count_day_of_sim_contract_end"] <= 30),
        },
        "last_synced":   sync["last_synced"].isoformat() if sync and sync["last_synced"] else None,
    }


# ── Device Management (write helpers) ──────────────────────────────

def _write_devmgmt_to_db(laptops, sims):
    """Shared write path used by both /sync (PA push) and /upload (manual UI).
    Replaces all rows (TRUNCATE + INSERT). Returns (laptops_count, sims_count).
    Currently PostgreSQL-only — writes are not wired for MSSQL."""
    conn = dbmod.getconn()
    try:
        with conn.cursor() as cur:
            cur.execute("TRUNCATE laptops")
            for r in laptops:
                cur.execute("""
                    INSERT INTO laptops (
                        serial_number, device_name, unique_device, enrollment_date,
                        manufacturer, model, category, primary_user_upn,
                        primary_user_display_name, device_per_user, remarks, status
                    ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                    ON CONFLICT (serial_number) DO NOTHING
                """, (
                    r.get("serial_number") or r.get("Serial number") or r.get("Serial Number"),
                    r.get("device_name") or r.get("Device name") or r.get("Device Name"),
                    r.get("unique_device") or r.get("Unique Device"),
                    r.get("enrollment_date") or r.get("Enrollment Date"),
                    r.get("manufacturer") or r.get("Manufacturer"),
                    r.get("model") or r.get("Model"),
                    r.get("category") or r.get("Category"),
                    r.get("primary_user_upn") or r.get("Primary user UPN") or r.get("Primary User UPN"),
                    r.get("primary_user_display_name") or r.get("Primary user display name") or r.get("Primary User Display Name"),
                    r.get("device_per_user") or r.get("Device per user"),
                    r.get("remarks") or r.get("Remarks"),
                    r.get("status") or r.get("Status"),
                ))

            cur.execute("TRUNCATE sim_cards")
            for r in sims:
                cur.execute("""
                    INSERT INTO sim_cards (
                        sim_no, telcom, contract_site, contract_no, monthly_fee,
                        sim_owner, contract_start_date, contract_end_date,
                        site_contract, count_day_of_sim_contract_end
                    ) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                    ON CONFLICT (sim_no) DO NOTHING
                """, (
                    str(r.get("sim_no") or r.get("SIM No") or r.get("SIM no")),
                    r.get("telcom") or r.get("Telcom"),
                    r.get("contract_site") or r.get("Contract Site"),
                    str(r.get("contract_no") or r.get("Contract No") or ""),
                    r.get("monthly_fee") or r.get("Monthly Fee"),
                    r.get("sim_owner") or r.get("SIM Owner"),
                    r.get("contract_start_date") or r.get("Contract Start Date"),
                    r.get("contract_end_date") or r.get("Contract End Date"),
                    r.get("site_contract") or r.get("Site Contract"),
                    r.get("count_day_of_sim_contract_end") or r.get("Count Day of SIM Contract End"),
                ))

            cur.execute("""
                UPDATE device_mgmt_sync
                SET last_synced=NOW(), laptops_count=%s, sims_count=%s
                WHERE id=1
            """, (len(laptops), len(sims)))
        conn.commit()
        log.info(f"Device Management synced: {len(laptops)} laptops, {len(sims)} sims")
        return len(laptops), len(sims)
    except Exception:
        conn.rollback()
        raise
    finally:
        dbmod.putconn(conn)


def _parse_excel_devmgmt(file_stream):
    """Parse the uploaded .xlsx workbook and return (laptops_list, sims_list).
    Expects two named sheets: 'Laptop' and 'SIM'."""
    import datetime as _dt
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("openpyxl not installed on server")

    wb = load_workbook(file_stream, read_only=True, data_only=True)
    if "Laptop" not in wb.sheetnames or "SIM" not in wb.sheetnames:
        raise ValueError(f"File must have 'Laptop' and 'SIM' sheets. Found: {wb.sheetnames}")

    def _val(v):
        if isinstance(v, (_dt.datetime, _dt.date)):
            return v.strftime("%Y-%m-%d")
        return v

    def _rows(sheet_name, wanted_headers=None):
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = rows[0]
        out = []
        for row in rows[1:]:
            if not row or len(row) < 2 or row[1] in (None, ""):   # skip rows without ID in col 2
                continue
            d = {}
            for i, v in enumerate(row):
                if i < len(headers) and headers[i]:
                    h = headers[i]
                    if wanted_headers is None or h in wanted_headers:
                        d[h] = _val(v)
            out.append(d)
        return out

    laptops = _rows("Laptop")  # all columns
    SIM_WANT = {"Telcom","SIM No","Contract Site","Contract No","Monthly Fee",
                "SIM Owner","Contract Start Date","Contract End Date",
                "Site Contract","Count Day of SIM Contract End"}
    sims = _rows("SIM", SIM_WANT)
    return laptops, sims


# ── Device Management (sync from Power Automate) ────────────────────

@app.route("/api/device-management/sync", methods=["POST"])
def api_devmgmt_sync():
    # Auth via Bearer token
    auth = request.headers.get("Authorization", "")
    if not SYNC_TOKEN or auth != f"Bearer {SYNC_TOKEN}":
        return jsonify({"ok": False, "message": "unauthorized"}), 401

    payload = request.get_json(silent=True) or {}
    laptops = payload.get("laptops") or []
    sims    = payload.get("sim_cards") or payload.get("sims") or []

    try:
        lc, sc = _write_devmgmt_to_db(laptops, sims)
        return jsonify({"ok": True, "laptops_count": lc, "sims_count": sc})
    except Exception as e:
        log.exception("device-management sync failed")
        return jsonify({"ok": False, "message": str(e)}), 500


# ── Device Management (manual Excel upload from dashboard UI) ───────

@app.route("/api/device-management/upload", methods=["POST"])
@acctmod.admin_required
def api_devmgmt_upload():
    f = request.files.get("file")
    if f is None or not f.filename:
        return jsonify({"ok": False, "message": "No file uploaded (expected multipart field 'file')"}), 400
    if not f.filename.lower().endswith(".xlsx"):
        return jsonify({"ok": False, "message": f"Only .xlsx supported. Got: {f.filename}"}), 400

    try:
        laptops, sims = _parse_excel_devmgmt(f.stream)
        lc, sc = _write_devmgmt_to_db(laptops, sims)
        return jsonify({"ok": True, "laptops_count": lc, "sims_count": sc, "filename": f.filename})
    except ValueError as e:
        return jsonify({"ok": False, "message": str(e)}), 400
    except Exception as e:
        log.exception("device-management upload failed")
        return jsonify({"ok": False, "message": str(e)}), 500


# ── HTTP routes ─────────────────────────────────────────────────────

@app.route("/")
@acctmod.login_required
def device_management():
    return render_template("device_management.html", user=acctmod.current_user())


@app.route("/api/device-management")
@acctmod.login_required
def api_devmgmt_read():
    try:
        user = acctmod.current_user()
        allowed = None if user["role"] == "admin" else set(user["sites"])
        return jsonify({"ok": True, "data": fetch_device_management(allowed)})
    except Exception as e:
        log.exception("device-management read failed")
        return jsonify({"ok": False, "message": str(e)}), 500


@app.route("/api/health")
def api_health():
    return jsonify({"ok": True, "ts": datetime.now(timezone.utc).isoformat()})


# ── Accounts: sign-in, session, self-service password change ───────

@app.route("/login", methods=["GET"])
def login_page():
    acctmod.ensure_schema()
    if acctmod.current_user():
        return redirect(url_for("device_management"))
    return render_template("login.html")


@app.route("/api/auth/login", methods=["POST"])
def api_auth_login():
    acctmod.ensure_schema()
    payload = request.get_json(silent=True) or {}
    user = acctmod.verify_login(payload.get("email"), payload.get("password"))
    if not user:
        return jsonify({"ok": False, "message": "Wrong email or password."}), 401
    acctmod.login_session(user)
    return jsonify({"ok": True, "user": user})


@app.route("/api/auth/logout", methods=["POST"])
def api_auth_logout():
    acctmod.logout_session()
    return jsonify({"ok": True})


@app.route("/api/auth/me")
def api_auth_me():
    user = acctmod.current_user()
    if not user:
        return jsonify({"ok": False, "message": "Not signed in"}), 401
    return jsonify({"ok": True, "user": user})


@app.route("/api/auth/change-password", methods=["POST"])
@acctmod.login_required
def api_auth_change_password():
    payload = request.get_json(silent=True) or {}
    user = acctmod.current_user()
    try:
        acctmod.set_own_password(user["email"], payload.get("current_password"), payload.get("new_password"))
        return jsonify({"ok": True})
    except ValueError as e:
        return jsonify({"ok": False, "message": str(e)}), 400


# ── Accounts: management page (admin-gated) ─────────────────────────

@app.route("/accounts")
@acctmod.admin_required
def accounts_page():
    return render_template("account_management.html", user=acctmod.current_user())


@app.route("/api/accounts")
@acctmod.admin_required
def api_accounts_list():
    return jsonify({"ok": True, "accounts": acctmod.list_accounts(), "sites": acctmod.SITES})


@app.route("/api/accounts", methods=["POST"])
@acctmod.admin_required
def api_accounts_create():
    payload = request.get_json(silent=True) or {}
    try:
        user = acctmod.create_account(
            payload.get("email"), payload.get("name"), payload.get("password"),
            payload.get("role") or "user", payload.get("sites") or [],
        )
        return jsonify({"ok": True, "account": user})
    except ValueError as e:
        return jsonify({"ok": False, "message": str(e)}), 400


@app.route("/api/accounts/<email>", methods=["PUT"])
@acctmod.admin_required
def api_accounts_update(email):
    payload = request.get_json(silent=True) or {}
    try:
        user = acctmod.update_account(
            email, payload.get("name"), payload.get("role") or "user",
            payload.get("active", True), payload.get("sites") or [],
            password=payload.get("password") or None,
        )
        return jsonify({"ok": True, "account": user})
    except ValueError as e:
        return jsonify({"ok": False, "message": str(e)}), 400


@app.route("/api/accounts/<email>", methods=["DELETE"])
@acctmod.admin_required
def api_accounts_delete(email):
    me = acctmod.current_user()
    if str(email).strip().lower() == me["email"]:
        return jsonify({"ok": False, "message": "You can't delete your own account."}), 400
    acctmod.delete_account(email)
    return jsonify({"ok": True})


# ── Configuration page (admin-gated) ────────────────────────────────

def _is_admin():
    return session.get("is_admin") is True


@app.route("/config/login", methods=["GET", "POST"])
def config_login():
    if request.method == "POST":
        pw = (request.form.get("password") or "").strip()
        if pw == ADMIN_PASSWORD:
            session["is_admin"] = True
            return redirect(url_for("config_page"))
        return render_template("db_config_login.html", error="Wrong password.")
    return render_template("db_config_login.html")


@app.route("/config/logout")
def config_logout():
    session.pop("is_admin", None)
    return redirect(url_for("config_login"))


@app.route("/config")
def config_page():
    if not _is_admin():
        return redirect(url_for("config_login"))
    cfg = dbmod.load_config()
    return render_template("db_config.html", cfg=cfg)


@app.route("/api/config/test", methods=["POST"])
def api_config_test():
    if not _is_admin():
        return jsonify({"ok": False, "message": "unauthorized"}), 401
    cfg = request.get_json(silent=True) or {}
    return jsonify(dbmod.test_connection(cfg))


@app.route("/api/config/save", methods=["POST"])
def api_config_save():
    if not _is_admin():
        return jsonify({"ok": False, "message": "unauthorized"}), 401
    cfg = request.get_json(silent=True) or {}
    # Sanity: run a test first to avoid saving broken config
    result = dbmod.test_connection(cfg)
    if not result.get("ok"):
        return jsonify({"ok": False, "message": f"Test failed — not saved. {result.get('message')}"}), 400
    try:
        dbmod.save_config(cfg)
        dbmod.reset_pool()
        return jsonify({"ok": True, "message": "Saved and pool rebuilt."})
    except Exception as e:
        log.exception("save config failed")
        return jsonify({"ok": False, "message": str(e)}), 500


@app.route("/api/config/reset", methods=["POST"])
def api_config_reset():
    if not _is_admin():
        return jsonify({"ok": False, "message": "unauthorized"}), 401
    try:
        dbmod.delete_config()
        dbmod.reset_pool()
        return jsonify({"ok": True, "message": "Reverted to DATABASE_URL env."})
    except Exception as e:
        log.exception("reset config failed")
        return jsonify({"ok": False, "message": str(e)}), 500


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8080"))
    log.info(f"Device Management starting on :{port}")
    app.run(host="0.0.0.0", port=port, debug=False, threaded=True)
